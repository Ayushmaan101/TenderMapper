"""Excel export (checklist 4.4).

Pure openpyxl workbook building - no Streamlit import, matching the
established pattern in this codebase (ocr/pipeline.py, search/bm25_index.py,
verify/groq_verifier.py are all Streamlit-free; the UI-facing sliver that
wires a download button lives in run/export_ui.py). This split is also
what makes the workbook itself directly unit-testable without AppTest.

One worksheet per configured table (schema_table), one data row per
configured column (schema_column) - the same row orientation
run/results.py settled on in checklist 4.1: config storage is uniform
regardless of how a table "looked" in the source tender document, so one
row per schema_column is what stays correct for an arbitrary, freely
reconfigured schema.

Data source: st.session_state["resolution_results"] (passed in by the
caller as `results_store`, never read directly here - keeps this module
Streamlit-free) - the REVIEWER-EDITED state, not a fresh re-run of
search+verify. A cell a human corrected via st.data_editor in the Run
tab is exactly what lands in the export, because run/results.py already
writes every edit back into that same dict (checklist 4.1's persistence
contract) before this module ever sees it.

Exported columns, exactly as specified: Status, Requirement, PDF Name,
Page Number / Range, Confidence, Match Snippet. Deliberately does NOT
include the short "Column" label (e.g. "Item 4") that the Run tab's grid
shows alongside Requirement - the full requirement text is what a reader
outside this app needs; the short internal label isn't.

Sheet name sanitization: Excel worksheet names must be <=31 characters,
must not contain \\ / ? * [ ] :, must not be blank, and must be unique
within the workbook (case-insensitively - Excel does not allow "Table A"
and "TABLE A" as two different sheets). All four are handled here, not
left to openpyxl (which raises on some of these itself, but not all, and
not with a graceful fallback).
"""
from __future__ import annotations

import io
import re
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from db import crud
from verify.groq_verifier import VerificationResult, is_flagged

EXPORT_HEADERS = ["Status", "Requirement", "PDF Name", "Page Number / Range", "Confidence", "Match Snippet"]
_COLUMN_WIDTHS = [12, 60, 24, 18, 12, 60]  # matches EXPORT_HEADERS order

DEFAULT_FILENAME = "tender_compliance_mapping.xlsx"

_BLANK_RESULT = VerificationResult(
    pdf_name="", page_number_or_range="", confidence=0.0, match_snippet="",
    reasoning="", success=False, error=None,
)

_INVALID_SHEET_CHARS_RE = re.compile(r"[\\/?*\[\]:]")
_INVALID_FILENAME_CHARS_RE = re.compile(r'[<>:"/\\|?*\x00-\x1f]')

MAX_SHEET_NAME_LENGTH = 31


def sanitize_sheet_name(name: str, existing_lower: set[str]) -> str:
    """Excel-legal, unique (case-insensitively) worksheet name. Mutates
    existing_lower by adding the returned name's lowercased form, so
    repeated calls across a workbook build correctly disambiguate.
    """
    cleaned = _INVALID_SHEET_CHARS_RE.sub("_", name).strip()
    if not cleaned:
        cleaned = "Sheet"
    cleaned = cleaned[:MAX_SHEET_NAME_LENGTH]

    base = cleaned
    candidate = cleaned
    suffix = 1
    while candidate.lower() in existing_lower:
        suffix_text = f" ({suffix})"
        candidate = base[: MAX_SHEET_NAME_LENGTH - len(suffix_text)] + suffix_text
        suffix += 1

    existing_lower.add(candidate.lower())
    return candidate


def sanitize_filename(company_name: str) -> str:
    """A safe .xlsx filename derived from the company name, falling back
    to DEFAULT_FILENAME for a blank/whitespace-only or fully-invalid name.
    """
    cleaned = _INVALID_FILENAME_CHARS_RE.sub("_", company_name or "").strip()
    cleaned = cleaned.strip(". ")  # Windows rejects filenames ending in a dot or space
    if not cleaned:
        return DEFAULT_FILENAME
    return f"{cleaned}.xlsx"


def _write_sheet(ws: Worksheet, columns, results_store: dict[int, VerificationResult]) -> None:
    ws.append(EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for col in columns:
        result = results_store.get(col.id, _BLANK_RESULT)
        status = "Flagged" if is_flagged(result) else "Resolved"
        ws.append([
            status,
            col.requirement_text,
            result.pdf_name,
            result.page_number_or_range,
            result.confidence,
            result.match_snippet,
        ])

    for idx, width in enumerate(_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    requirement_col = EXPORT_HEADERS.index("Requirement") + 1
    snippet_col = EXPORT_HEADERS.index("Match Snippet") + 1
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[requirement_col - 1].alignment = Alignment(wrap_text=True, vertical="top")
        row[snippet_col - 1].alignment = Alignment(wrap_text=True, vertical="top")
        row[EXPORT_HEADERS.index("Confidence")].number_format = "0.00"


def build_workbook(conn, results_store: Optional[dict[int, VerificationResult]] = None) -> Workbook:
    """Build the full export workbook from the CURRENT config (conn) and
    the reviewer-edited results_store. results_store defaults to {} -
    every row then exports as an unresolved/Flagged placeholder, which is
    correct (not an error) for a company that's never been processed.
    """
    results_store = results_store or {}

    wb = Workbook()
    wb.remove(wb.active)  # drop openpyxl's default "Sheet" - we add our own per table

    tables = crud.get_tables(conn)
    existing_sheet_names_lower: set[str] = set()

    for table in tables:
        columns = crud.get_columns(conn, table.id)
        sheet_name = sanitize_sheet_name(table.name, existing_sheet_names_lower)
        ws = wb.create_sheet(title=sheet_name)
        _write_sheet(ws, columns, results_store)

    if not wb.sheetnames:
        # A workbook must have at least one sheet - a zero-table config
        # would otherwise produce an unopenable file.
        ws = wb.create_sheet(title="No Tables Configured")
        ws.append(["No tables are configured yet."])

    return wb


def workbook_to_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_to_bytes(conn, results_store: Optional[dict[int, VerificationResult]] = None) -> bytes:
    """Convenience one-shot: build_workbook() + workbook_to_bytes()."""
    return workbook_to_bytes(build_workbook(conn, results_store))
